# -*- coding: utf-8 -*-
"""Gera tabela Markdown + XLSX de fellowships médicos a partir de JSONs por país.

Uso: python build_fellowships.py [dataset_dir]
  dataset_dir contém <cc>-fellowships.json (br, ca, uk, au... ou quaisquer chaves).
  Ajuste FILES/COUNTRY_ORDER abaixo para outros países/especialidades.
"""
import json, re, sys, pathlib

BASE = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else pathlib.Path.home() / "Desktop" / "fellowships-coluna"
OUT_MD = BASE / "tabela-fellowships-coluna.md"
OUT_XLSX = pathlib.Path(str(BASE) + "-2027-2028.xlsx")

# cotas: 1 USD = X (atualize via: curl -s https://open.er-api.com/v6/latest/USD)
FX = {"USD": 1.0, "BRL": 5.222334, "CAD": 1.424002, "GBP": 0.756313, "AUD": 1.438733}

FILES = [("br", "br-fellowships.json"), ("ca", "ca-fellowships.json"),
         ("uk", "uk-fellowships.json"), ("au", "au-fellowships.json")]
FILE_COUNTRY = {"br": "Brasil", "ca": "Canadá", "uk": "Reino Unido (Inglaterra e Escócia)", "au": "Austrália"}
COUNTRY_ORDER = {"Brasil": 0, "Canadá": 1, "Reino Unido (Inglaterra e Escócia)": 2, "Austrália": 3}

def clean_cell(s):
    if isinstance(s, (dict, list)):
        s = json.dumps(s, ensure_ascii=False)
    return re.sub(r"\s+", " ", (s or "").replace("|", "/").replace("\n", " ")).strip()

def parse_num(s):
    s = s.strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None

MONEY_PAT = [("CAD", r"CAD\s*\$?\s*([\d.,]+)"), ("AUD", r"AUD\s*\$?\s*([\d.,]+)"),
             ("GBP", r"£\s*([\d.,]+)"), ("BRL", r"R\$\s*([\d.,]+)")]

def usd_text(comp):
    comp = clean_cell(comp)
    if not comp:
        return ""
    lo = comp.lower()
    if lo.startswith(("não divulgado", "nao divulgado")):
        return "não divulgado"
    if lo.startswith(("não encontrado", "nao encontrado")):
        return "não encontrado"
    if lo.startswith(("sem remuneração", "sem remuneracao", "não remunerado")):
        return "sem remuneração"
    for ccy, pat in MONEY_PAT:
        m = re.search(pat, comp)
        if not m:
            continue
        v = parse_num(m.group(1))
        if v is None:
            continue
        mult = 12.0 if ("mês" in lo or "mensal" in lo) and "quinzena" not in lo else (26.0 if "quinzena" in lo else 1.0)
        usd = v * mult / FX[ccy]
        return f"≈ US$ {usd/1000:.0f}k/ano" if usd >= 10000 else f"≈ US$ {usd:,.0f}/ano"
    return ""

def short_comp(comp, usd):
    if usd in ("não divulgado", "não encontrado", "sem remuneração"):
        return usd
    base = re.sub(r"\s*\(as of 2018\)|\(valor de 2018\)", "", comp or "", flags=re.I)
    base = clean_cell(base)
    if usd:
        return (base[:70] + "… " + usd) if len(base) > 72 else (base + " " + usd)
    return (base[:95] + "…") if len(base) > 97 else base

def category(p):
    t = (p.get("type") or "").lower().strip()
    if t.startswith(("diretório", "diretorio", "directory", "portal")):
        return "diretorio"
    if t.startswith(("research", "observership", "observer", "clinical visit")) or "research fellowship" in t[:40]:
        return "obs/pesquisa"
    return "clinico"

def accepts_target(p):
    v = p.get("accepts_neurosurgeons") or p.get("accepts_target")
    if isinstance(v, dict):
        a = v.get("accepts")
        return "sim" if a is True else ("nao" if a in (False, None) else clean_cell(a).lower()[:3] or "nao")
    a = clean_cell(v).lower()
    if a.startswith("sim"):
        return "sim"
    if a.startswith("parcial"):
        return "parcial"
    return "nao"

def foco_tags(p):
    f = (p.get("focus") or "").lower()
    tags = []
    if any(k in f for k in ["minimamente invasiv", " mis", "mis ", "endoscop", "biportal", "percutân", "percutan"]): tags.append("MIS/endoscopia")
    if any(k in f for k in ["deformidade", "escoliose"]): tags.append("deformidade")
    if any(k in f for k in ["tumor", "neoplás", "neoplas", "oncol", "metastát", "metastat"]): tags.append("oncologia")
    if "trauma" in f: tags.append("trauma")
    if any(k in f for k in ["pediátric", "pediatric", "criança", "crianca"]): tags.append("pediátrica")
    if any(k in f for k in ["intradural", "intradurais", "medular", "craniovertebral", "cranio-cervical"]): tags.append("intradural/cranio-cervical")
    if "cervical" in f: tags.append("cervical")
    if any(k in f for k in ["degenerativ", "lombar"]): tags.append("degenerativa")
    return ", ".join(tags[:4]) if tags else "geral"

CONTACT_OVERRIDE = {"bristol spinal": "Mr Mark Nowell (contact pelo e-mail do programa)"}

def contato(p):
    key = (p.get("program", "") or "").lower()
    for k, v in CONTACT_OVERRIDE.items():
        if k in key:
            return v
    h = (p.get("head_of_spine") or "").strip()
    if isinstance(h, dict):
        h = h.get("name") or json.dumps(h, ensure_ascii=False)
    if not h or h.lower().startswith(("não encontrado", "nao encontrado")):
        return "não encontrado"
    for sep in [" — ", " – ", "; ", " ("]:
        if sep in h:
            h = h.split(sep)[0]
    return clean_cell(h)[:60]

def entrada(p):
    parts = []
    for f in ["start_availability", "application_window"]:
        v = (p.get(f) or "").strip()
        if not v:
            continue
        v = clean_cell(v)
        lv = v.lower()
        if lv.startswith(("não divulgado", "nao divulgado", "não encontrado", "nao encontrado")):
            rest = v.split("—", 1)[-1].strip() if "—" in v else ""
            if rest and not rest.lower().startswith("não") and len(rest) > 5:
                v = rest
            else:
                continue
        parts.append(v)
    s = " | ".join(dict.fromkeys(parts))
    if not s:
        return "não divulgado"
    return clean_cell(s)[:118] + ("…" if len(s) > 120 else "")

# grade de acesso: chave (substring do programa/instituição) -> fácil | médio | difícil
DIFF = {}

def acesso(p):
    k = ((p.get("program") or "") + " " + (p.get("institution") or "")).lower()
    for key, val in DIFF.items():
        if key in k:
            return val
    return "fácil" if p.get("_country") == "Brasil" else "médio"

def load_all():
    rows = []
    for key, fname in FILES:
        f = BASE / fname
        if not f.exists():
            print(f"[aviso] {fname} ausente — pulando")
            continue
        obj = json.loads(f.read_text(encoding="utf-8"))
        plist, country = (obj, FILE_COUNTRY[key]) if isinstance(obj, list) else (obj.get("programs", []), obj.get("country", FILE_COUNTRY[key]))
        for p in plist:
            p = dict(p)
            p["_country"] = country
            rows.append(p)
    return rows

def build():
    rows = load_all()
    main, secondary = [], []
    for p in rows:
        cat = category(p)
        an = accepts_target(p)
        usd = usd_text(p.get("compensation", ""))
        rec = {
            "pais": p["_country"], "program": clean_cell(p.get("program", "")),
            "institution": clean_cell(p.get("institution", "")), "city": clean_cell(p.get("city", "")),
            "contato": contato(p), "entrada": entrada(p),
            "email": clean_cell(p.get("contact_email", "")) or "não encontrado",
            "site": (p.get("website") or "").strip(), "comp": clean_cell(p.get("compensation", "")),
            "usd": usd, "comp_short": short_comp(p.get("compensation", ""), usd),
            "foco": foco_tags(p), "acesso": acesso(p), "cat": cat, "aceita": an,
            "type": clean_cell(p.get("type", "")), "lic": clean_cell(p.get("licensing_requirements", "")),
            "head_source": (p.get("head_source") or "").strip(),
            "start": clean_cell(p.get("start_availability", "")), "window": clean_cell(p.get("application_window", "")),
            "sources": " ; ".join(p.get("sources", []) or ([p["website"]] if p.get("website") else [])),
        }
        if cat == "diretorio":
            continue
        (secondary if (cat == "obs/pesquisa" or an == "nao") else main).append(rec)

    rank = {"fácil": 0, "médio": 1, "difícil": 2}
    for lst in (main, secondary):
        lst.sort(key=lambda r: (COUNTRY_ORDER.get(r["pais"], 9), rank.get(r["acesso"], 3), r["program"]))

    md = ["# Fellowships — tabela consolidada", "", "_Valores convertidos pela cotação USD do dia; dados não publicados = 'não encontrado'/'não divulgado'._", "", "## Fellowships clínicos formais", ""]
    cols = ["Programa", "Local", "Contato (diretor do fellowship)", "Entrada", "E-mail", "Site", "Remuneração", "Foco", "Acesso"]
    md.append("| " + " | ".join(cols) + " |")
    md.append("|" + "---|" * len(cols))
    for r in main:
        site = f"[site]({r['site']})" if r["site"] else "—"
        md.append("| " + " | ".join([r["program"][:90], f"{r['institution'][:55]} ({r['city'][:25]})", r["contato"], r["entrada"], r["email"][:60], site, r["comp_short"], r["foco"], r["acesso"] + (" ⚑" if r["aceita"] == "parcial" else "")]) + " |")
    md += ["", "## Observerships, research e aceitação não confirmada", ""]
    cols2 = ["Programa", "Local", "Tipo", "Contato", "E-mail", "Site", "Remuneração", "Acesso"]
    md.append("| " + " | ".join(cols2) + " |")
    md.append("|" + "---|" * len(cols2))
    for r in secondary:
        site = f"[site]({r['site']})" if r["site"] else "—"
        md.append("| " + " | ".join([r["program"][:90], f"{r['institution'][:55]} ({r['city'][:25]})", r["type"][:60], r["contato"], r["email"][:60], site, r["comp_short"], r["acesso"]]) + " |")
    OUT_MD.write_text("\n".join(md), encoding="utf-8")

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Fellowships"
    headers = ["País", "Categoria", "Programa", "Instituição", "Cidade", "Tipo", "Aceita neurocirurgiões",
               "Foco", "Contato (diretor/chefe)", "Fonte do contato", "Disponibilidade de entrada",
               "Janela de candidatura", "E-mail", "Site", "Remuneração (original)", "Remuneração (USD aprox.)",
               "Licenciamento/visto", "Acesso", "Fontes"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    allrows = main + secondary
    allrows.sort(key=lambda r: (COUNTRY_ORDER.get(r["pais"], 9), r["cat"], rank.get(r["acesso"], 3), r["program"]))
    for r in allrows:
        ws.append([r["pais"], r["cat"], r["program"], r["institution"], r["city"], r["type"],
                   {"sim": "sim", "parcial": "parcial (⚑)", "nao": "não confirmado"}[r["aceita"]],
                   r["foco"], r["contato"], r["head_source"], r["start"], r["window"], r["email"],
                   r["site"], r["comp"], r["usd"], r["lic"], r["acesso"], r["sources"]])
    widths = [14, 12, 45, 40, 18, 30, 14, 28, 30, 35, 35, 35, 32, 40, 35, 18, 45, 9, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A" + chr(64 + i - 26)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(OUT_XLSX)
    print(f"main={len(main)} secondary={len(secondary)} total={len(allrows)}")
    print("MD:", OUT_MD)
    print("XLSX:", OUT_XLSX)

if __name__ == "__main__":
    build()
